"""
导出路由
提供 /chat/api/export 端点
"""

import io
import logging
from fastapi import APIRouter, Depends, HTTPException
from fastapi.responses import StreamingResponse
from sqlalchemy.orm import Session
from pydantic import BaseModel

from app.core.database import get_db
from app.core.security import get_current_user
from app.models.models import Dialog, DialogMessage

logger = logging.getLogger(__name__)

router = APIRouter(prefix="/export", tags=["导出功能"])


class ExportRequest(BaseModel):
    dialog_id: str
    format: str = "txt"  # txt/markdown/docx


@router.post("")
async def export_script(
    data: ExportRequest,
    db: Session = Depends(get_db),
    current_user: dict = Depends(get_current_user),
):
    """导出对话中的脚本"""
    dialog = db.query(Dialog).filter(Dialog.id == data.dialog_id).first()
    if not dialog:
        raise HTTPException(status_code=404, detail="对话不存在")
    if dialog.user_id != current_user["user_id"]:
        raise HTTPException(status_code=403, detail="无权导出此对话")

    # 获取对话的最后一条AI消息
    last_ai_message = db.query(DialogMessage).filter(
        DialogMessage.dialog_id == data.dialog_id,
        DialogMessage.role == "assistant"
    ).order_by(DialogMessage.created_at.desc()).first()
    
    if not last_ai_message:
        raise HTTPException(status_code=404, detail="未找到可导出的脚本内容")
    
    content = last_ai_message.content
    
    if data.format == "txt":
        return _export_as_txt(content)
    elif data.format == "markdown":
        return _export_as_markdown(content)
    elif data.format == "docx":
        return _export_as_docx(content)
    else:
        raise HTTPException(status_code=400, detail="不支持的导出格式")


def _export_as_txt(content: str) -> StreamingResponse:
    """导出为纯文本"""
    # 移除Markdown格式
    import re
    text = content
    # 移除标题标记
    text = re.sub(r'^#{1,6}\s+', '', text, flags=re.MULTILINE)
    # 移除加粗标记
    text = re.sub(r'\*\*(.*?)\*\*', r'\1', text)
    # 移除斜体标记
    text = re.sub(r'\*(.*?)\*', r'\1', text)
    # 移除列表标记
    text = re.sub(r'^[-*]\s+', '', text, flags=re.MULTILINE)
    text = re.sub(r'^\d+\.\s+', '', text, flags=re.MULTILINE)
    # 移除分隔线
    text = re.sub(r'^---+$', '', text, flags=re.MULTILINE)
    # 清理多余空行
    text = re.sub(r'\n{3,}', '\n\n', text)
    
    buffer = io.BytesIO()
    buffer.write(text.encode('utf-8'))
    buffer.seek(0)
    
    return StreamingResponse(
        buffer,
        media_type="text/plain; charset=utf-8",
        headers={"Content-Disposition": "attachment; filename=script.txt"}
    )


def _export_as_markdown(content: str) -> StreamingResponse:
    """导出为Markdown"""
    buffer = io.BytesIO()
    buffer.write(content.encode('utf-8'))
    buffer.seek(0)
    
    return StreamingResponse(
        buffer,
        media_type="text/markdown; charset=utf-8",
        headers={"Content-Disposition": "attachment; filename=script.md"}
    )


def _export_as_docx(content: str) -> StreamingResponse:
    """导出为Word文档"""
    try:
        from docx import Document
        from docx.shared import Pt, Inches
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        import re
        
        doc = Document()
        
        # 设置默认字体
        style = doc.styles['Normal']
        font = style.font
        font.name = 'Microsoft YaHei'
        font.size = Pt(11)
        
        lines = content.split('\n')
        
        for line in lines:
            line = line.strip()
            if not line:
                doc.add_paragraph('')
                continue
            
            # 处理标题
            if line.startswith('# '):
                heading = doc.add_heading(line[2:].strip(), level=1)
            elif line.startswith('## '):
                heading = doc.add_heading(line[3:].strip(), level=2)
            elif line.startswith('### '):
                heading = doc.add_heading(line[4:].strip(), level=3)
            elif line.startswith('#### '):
                heading = doc.add_heading(line[5:].strip(), level=4)
            elif line.startswith('---'):
                doc.add_paragraph('─' * 50)
            elif line.startswith('- ') or line.startswith('* '):
                # 列表项
                text = line[2:].strip()
                text = re.sub(r'\*\*(.*?)\*\*', r'\1', text)
                doc.add_paragraph(text, style='List Bullet')
            elif re.match(r'^\d+\.\s', line):
                # 有序列表
                text = re.sub(r'^\d+\.\s', '', line)
                text = re.sub(r'\*\*(.*?)\*\*', r'\1', text)
                doc.add_paragraph(text, style='List Number')
            else:
                # 普通段落
                text = line
                text = re.sub(r'\*\*(.*?)\*\*', r'\1', text)
                text = re.sub(r'\*(.*?)\*', r'\1', text)
                doc.add_paragraph(text)
        
        buffer = io.BytesIO()
        doc.save(buffer)
        buffer.seek(0)
        
        return StreamingResponse(
            buffer,
            media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            headers={"Content-Disposition": "attachment; filename=script.docx"}
        )
    except ImportError:
        # 如果没有python-docx，降级为txt
        logger.warning("python-docx未安装，降级为txt格式导出")
        return _export_as_txt(content)
