from fastapi import APIRouter, Depends, HTTPException, UploadFile, File
from fastapi.responses import StreamingResponse
from sqlalchemy.orm import Session
from typing import List
import openpyxl
import io

from app.core.database import get_db
from app.models.models import ForbiddenWord
from app.api.admin import get_current_admin
from app.schemas.schemas import (
    ForbiddenWordCreate,
    ForbiddenWordUpdate,
    ForbiddenWordResponse,
    ForbiddenWordImportItem,
)
from app.services.forbidden_word_service import forbidden_word_service

router = APIRouter(prefix="/api/forbidden-words", tags=["forbidden-words"])


@router.get("/public")
async def get_public_forbidden_words(db: Session = Depends(get_db)):
    """获取违禁词列表（公开接口，用于前端过滤）"""
    words = db.query(ForbiddenWord).order_by(ForbiddenWord.created_at.desc()).all()
    return [{"id": w.id, "word": w.word, "replacement": w.replacement, "level": w.level} for w in words]


@router.get("", response_model=List[ForbiddenWordResponse])
async def get_all_forbidden_words(db: Session = Depends(get_db), _admin=Depends(get_current_admin)):
    """获取所有违禁词（管理员接口）"""
    words = db.query(ForbiddenWord).order_by(ForbiddenWord.created_at.desc()).all()
    return words


def _auto_level(word: str) -> str:
    """根据字数自动判断级别"""
    return "single" if len(word) == 1 else "multi"


@router.post("", response_model=ForbiddenWordResponse)
async def create_forbidden_word(
    data: ForbiddenWordCreate,
    db: Session = Depends(get_db),
    _admin=Depends(get_current_admin)
):
    """创建违禁词"""
    existing = db.query(ForbiddenWord).filter(ForbiddenWord.word == data.word).first()
    if existing:
        raise HTTPException(status_code=400, detail=f"违禁词 '{data.word}' 已存在")

    level = data.level if data.level else _auto_level(data.word)
    word = ForbiddenWord(word=data.word, replacement=data.replacement, level=level)
    db.add(word)
    db.commit()
    db.refresh(word)
    forbidden_word_service.invalidate_cache()
    return word


@router.put("/{word_id}", response_model=ForbiddenWordResponse)
async def update_forbidden_word(
    word_id: int,
    data: ForbiddenWordUpdate,
    db: Session = Depends(get_db),
    _admin=Depends(get_current_admin)
):
    """更新违禁词"""
    word = db.query(ForbiddenWord).filter(ForbiddenWord.id == word_id).first()
    if not word:
        raise HTTPException(status_code=404, detail="违禁词不存在")

    if data.word is not None:
        existing = (
            db.query(ForbiddenWord)
            .filter(ForbiddenWord.word == data.word, ForbiddenWord.id != word_id)
            .first()
        )
        if existing:
            raise HTTPException(status_code=400, detail=f"违禁词 '{data.word}' 已存在")
        word.word = data.word

    if data.replacement is not None:
        word.replacement = data.replacement
    if data.level is not None:
        word.level = data.level or _auto_level(word.word)

    db.commit()
    db.refresh(word)
    forbidden_word_service.invalidate_cache()
    return word


@router.delete("/{word_id}")
async def delete_forbidden_word(
    word_id: int,
    db: Session = Depends(get_db),
    _admin=Depends(get_current_admin)
):
    """删除违禁词"""
    word = db.query(ForbiddenWord).filter(ForbiddenWord.id == word_id).first()
    if not word:
        raise HTTPException(status_code=404, detail="违禁词不存在")

    db.delete(word)
    db.commit()
    forbidden_word_service.invalidate_cache()
    return {"success": True}


@router.post("/import")
async def import_forbidden_words(
    file: UploadFile = File(...),
    db: Session = Depends(get_db),
    _admin=Depends(get_current_admin)
):
    """从Excel文件导入违禁词"""
    if not file.filename.endswith((".xlsx", ".xls")):
        raise HTTPException(status_code=400, detail="请上传Excel文件(.xlsx或.xls)")

    try:
        content = await file.read()
        wb = openpyxl.load_workbook(io.BytesIO(content))
        ws = wb.active

        imported_count = 0
        skipped_count = 0
        errors = []

        for row_idx, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
            if not row or not row[0]:
                continue

            word = str(row[0]).strip()
            replacement = str(row[1]).strip() if len(row) > 1 and row[1] else "***"
            level = str(row[2]).strip() if len(row) > 2 and row[2] else _auto_level(word)

            if not word:
                continue

            existing = (
                db.query(ForbiddenWord).filter(ForbiddenWord.word == word).first()
            )
            if existing:
                skipped_count += 1
                continue

            db.add(ForbiddenWord(word=word, replacement=replacement, level=level))
            imported_count += 1

        db.commit()
        forbidden_word_service.invalidate_cache()

        return {
            "success": True,
            "imported": imported_count,
            "skipped": skipped_count,
            "errors": errors,
        }

    except Exception as e:
        raise HTTPException(status_code=400, detail=f"导入失败: {str(e)}")


@router.delete("/batch/all")
async def delete_all_forbidden_words(
    db: Session = Depends(get_db),
    _admin=Depends(get_current_admin)
):
    """清空所有违禁词"""
    count = db.query(ForbiddenWord).delete()
    db.commit()
    forbidden_word_service.invalidate_cache()
    return {"success": True, "deleted": count}


@router.get("/template")
async def download_template(_admin=Depends(get_current_admin)):
    """下载违禁词导入模板"""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "违禁词模板"

    # 设置表头
    ws.append(["违禁词", "替换词", "级别"])

    # 添加示例数据
    ws.append(["示例违禁词1", "***", "multi"])
    ws.append(["示例违禁词2", "xxx", "multi"])
    ws.append(["单", "*", "single"])

    # 设置列宽
    ws.column_dimensions['A'].width = 20
    ws.column_dimensions['B'].width = 15
    ws.column_dimensions['C'].width = 10

    # 保存到内存
    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)

    return StreamingResponse(
        buffer,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=forbidden_words_template.xlsx"}
    )


@router.get("/check")
async def check_text_for_forbidden_words(
    text: str, db: Session = Depends(get_db), _admin=Depends(get_current_admin)
):
    """检测文本中的违禁词并返回替换后的文本"""
    return forbidden_word_service.process_text(text, db)
